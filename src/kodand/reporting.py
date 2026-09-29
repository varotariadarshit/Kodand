import base64
import html
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image, Paragraph, Preformatted, SimpleDocTemplate, Spacer

from kodand.categories import TAXONOMIES
from kodand.categories import category_color as _cat_color
from kodand.categories import taxonomy_of as _taxonomy_of


class _CatColorMap:
    def get(self, key, default="#888"):
        try:
            return _cat_color(key)
        except Exception:
            return default
    def __getitem__(self, key):
        return self.get(key)


_ASSETS_DIR = Path(__file__).parent / "webui" / "static"
_ICON_PATH = _ASSETS_DIR / "kodand-icon.png"


def _asset_b64(name: str) -> str:
    return base64.b64encode((_ASSETS_DIR / name).read_bytes()).decode()


_LOGO_DATA_URI = f"data:image/png;base64,{_asset_b64('kodand-logo-full.png')}"
_FAVICON_DATA_URI = f"data:image/png;base64,{_asset_b64('favicon-32.png')}"

CATEGORY_COLORS = _CatColorMap()
SEV_COLORS = {"critical": "#f43f5e", "high": "#fb923c", "medium": "#facc15",
              "low": "#38bdf8", "informational": "#94a3b8"}

_CSS = """
:root{--bg0:#05060a;--bg1:#0b1020;--glass:rgba(255,255,255,.06);--glass-strong:rgba(255,255,255,.10);--border:rgba(255,255,255,.14);--txt:#e8ecf6;--muted:#9aa4bd;--shadow:0 8px 40px rgba(0,0,0,.45);color-scheme:dark}
*{box-sizing:border-box}html{scroll-behavior:smooth}
body{margin:0;color:var(--txt);font:15px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;background:radial-gradient(1200px 700px at 12% -10%,rgba(99,102,241,.22),transparent 60%),radial-gradient(1000px 600px at 110% 10%,rgba(236,72,153,.16),transparent 55%),radial-gradient(900px 700px at 50% 120%,rgba(16,185,129,.14),transparent 55%),linear-gradient(160deg,var(--bg1),var(--bg0));background-attachment:fixed;min-height:100vh}
.wrap{max-width:1120px;margin:0 auto;padding:38px 22px 80px}
a{color:#8ab4ff;text-decoration:none}a:hover{text-decoration:underline}
code{font-family:"JetBrains Mono",ui-monospace,SFMono-Regular,Menlo,Consolas,monospace;background:rgba(255,255,255,.08);padding:2px 6px;border-radius:6px;font-size:.86em;color:#c9d4ff;word-break:break-all}
h1,h2,h3,h4{line-height:1.25;margin:0 0 .5em}
.glass{background:var(--glass);backdrop-filter:blur(16px) saturate(140%);-webkit-backdrop-filter:blur(16px) saturate(140%);border:1px solid var(--border);border-radius:18px;box-shadow:var(--shadow);padding:22px 24px;margin:18px 0}
.section-title{font-size:.82rem;letter-spacing:.18em;text-transform:uppercase;color:var(--muted);margin:26px 4px 6px}
.hero{position:relative;overflow:hidden}
.brand{display:flex;align-items:center;gap:18px;flex-wrap:wrap}
.brand img.logo-lockup{height:80px;width:auto;max-width:160px;flex:none}
.brand-label{font-weight:700;font-size:1.15rem}
.report-nav{display:flex;gap:10px;margin:0 0 4px}
.report-nav a{display:inline-flex;align-items:center;gap:6px;font-size:.82rem;font-weight:600;color:var(--txt);background:var(--glass-strong);border:1px solid var(--border);padding:8px 14px;border-radius:10px;text-decoration:none}
.report-nav a:hover{border-color:#8ab4ff;text-decoration:none}
.tagline{color:var(--muted);margin-top:2px}
.chips{display:flex;gap:10px;flex-wrap:wrap;margin-top:16px}
.chip{display:inline-flex;align-items:center;gap:8px;padding:8px 14px;border-radius:999px;background:var(--glass-strong);border:1px solid var(--border);font-size:.86rem}
.chip b{color:#fff}.dot{width:9px;height:9px;border-radius:50%}
.stats{display:grid;grid-template-columns:repeat(5,1fr);gap:14px}
.stat{background:var(--glass);border:1px solid var(--border);border-radius:14px;padding:16px;text-align:center;position:relative;overflow:hidden}
.stat .n{font-size:2rem;font-weight:800;line-height:1}.stat .l{font-size:.72rem;letter-spacing:.12em;text-transform:uppercase;color:var(--muted);margin-top:6px}
.stat::before{content:"";position:absolute;inset:0 0 auto 0;height:4px;background:var(--accent,#888)}
.charts{display:flex;flex-direction:column;align-items:center;gap:22px}
.donut-wrap{display:flex;flex-direction:column;align-items:center;justify-content:center;gap:14px}
.donut{width:190px;height:190px;border-radius:50%;position:relative;background:conic-gradient(var(--gc))}
.donut::after{content:"";position:absolute;inset:26px;border-radius:50%;background:linear-gradient(160deg,#0c1226,#070a14);box-shadow:inset 0 0 24px rgba(0,0,0,.6)}
.donut-center{position:absolute;inset:0;display:flex;flex-direction:column;align-items:center;justify-content:center;z-index:2}
.donut-center .big{font-size:1.8rem;font-weight:800}.donut-center .sm{font-size:.7rem;letter-spacing:.14em;text-transform:uppercase;color:var(--muted)}
.legend{display:flex;flex-direction:row;flex-wrap:wrap;justify-content:center;row-gap:8px}
.legend .row{display:flex;align-items:center;gap:7px;font-size:.86rem;white-space:nowrap;padding:0 16px}
.legend .row+.row{border-left:1px solid rgba(255,255,255,.14)}
.legend .row .swatch{width:10px;height:10px;border-radius:3px;flex:none}.legend .row .val{font-weight:700;color:var(--txt)}
.bars-cols{display:grid;grid-template-columns:repeat(3,1fr);gap:14px 26px;width:100%}
@media(max-width:900px){.bars-cols{grid-template-columns:1fr}.stats{grid-template-columns:repeat(2,1fr)}}
.bars{display:flex;flex-direction:column;gap:10px;justify-content:center;min-width:0}
.bar-row{display:grid;grid-template-columns:56px 1fr 26px;align-items:center;gap:8px;font-size:.82rem}
.bar-track{height:14px;border-radius:999px;background:rgba(255,255,255,.07);overflow:hidden}
.bar-fill{height:100%;width:0;border-radius:999px;transition:width 1.1s cubic-bezier(.2,.8,.2,1)}
.bar-row .cat{color:var(--muted)}.bar-row .cnt{text-align:right;font-variant-numeric:tabular-nums}
table{width:100%;border-collapse:separate;border-spacing:0;overflow:hidden;border-radius:12px}
th,td{padding:11px 14px;text-align:left;border-bottom:1px solid rgba(255,255,255,.08);vertical-align:top}
thead th{background:rgba(255,255,255,.06);font-size:.74rem;letter-spacing:.1em;text-transform:uppercase;color:var(--muted)}
tbody tr:hover{background:rgba(255,255,255,.03)}td ul{margin:0;padding-left:18px}td li{margin:2px 0}
.pill{display:inline-block;padding:3px 10px;border-radius:999px;font-size:.74rem;font-weight:600;border:1px solid transparent}
.scorebar{height:8px;border-radius:999px;background:rgba(255,255,255,.08);overflow:hidden;min-width:90px}
.scorebar i{display:block;height:100%;border-radius:999px}
.filters{display:flex;gap:8px;flex-wrap:wrap;margin:6px 4px 0}
.fbtn{cursor:pointer;padding:7px 14px;border-radius:999px;background:var(--glass-strong);border:1px solid var(--border);color:var(--txt);font-size:.82rem;transition:.2s}
.fbtn:hover{border-color:#8ab4ff}.fbtn.active{background:#8ab4ff;color:#0a1020;border-color:#8ab4ff;font-weight:700}
.finding{border-left:5px solid var(--sev,#888);position:relative}
.finding.resolved{opacity:.55}
.finding.resolved:hover{opacity:.85}
.finding .fhead{display:flex;align-items:flex-start;gap:12px;cursor:pointer}
.finding .ftitle{font-size:1.12rem;font-weight:700}.finding .fmeta{color:var(--muted);font-size:.82rem;margin-top:4px}
.finding .badges{margin-left:auto;display:flex;gap:8px;flex-shrink:0;flex-wrap:wrap;justify-content:flex-end}
.finding .fbody{margin-top:14px;display:grid;grid-template-columns:1fr 1fr;gap:14px 26px}
.finding .fbody.collapsed{display:none}
.finding h4{font-size:.74rem;letter-spacing:.1em;text-transform:uppercase;color:var(--muted);margin:0 0 4px}
.finding .full{grid-column:1/-1}
.ev{background:rgba(0,0,0,.28);border:1px solid rgba(255,255,255,.08);border-radius:10px;padding:10px 12px;margin:8px 0;font-size:.86rem}
.ev .hash{color:var(--muted);font-size:.72rem;word-break:break-all}
.caret{transition:transform .2s;color:var(--muted)}.finding.open .caret{transform:rotate(90deg)}
.toc{columns:2;gap:26px}@media(max-width:640px){.toc{columns:1}}.toc li{margin:4px 0;break-inside:avoid}
.toc-wrap{display:flex;flex-direction:column;gap:18px}
.toc-group h4{font-size:.74rem;letter-spacing:.1em;text-transform:uppercase;color:var(--muted);margin:0 0 8px}
.open-file-btn{cursor:pointer;border:1px solid rgba(255,255,255,.16);background:rgba(255,255,255,.06);color:#c9d4ff;font:inherit;font-size:.72rem;font-weight:600;padding:3px 10px;border-radius:20px;white-space:nowrap}
.open-file-btn:hover{background:rgba(138,180,255,.18);border-color:#8ab4ff}
.foot{color:var(--muted);font-size:.82rem;text-align:center;margin-top:30px}
.top-link{position:fixed;right:18px;bottom:18px;background:var(--glass-strong);border:1px solid var(--border);backdrop-filter:blur(10px);border-radius:12px;padding:10px 14px;font-size:.8rem}
.tag{display:inline-block;padding:2px 8px;border-radius:6px;font-size:.7rem;background:rgba(255,255,255,.08);color:#c9d4ff;margin:0 4px 4px 0}
.review-panel .row{display:flex;flex-wrap:wrap;align-items:center;gap:10px}
.review-panel input.rv-comment-input,.review-panel input.rv-chat-input{flex:1;min-width:200px;background:rgba(0,0,0,.28);border:1px solid rgba(255,255,255,.14);border-radius:8px;color:var(--txt);padding:8px 10px;font:inherit}
.review-panel button{cursor:pointer;border:1px solid #8ab4ff;background:#8ab4ff;color:#08101f;font-weight:700;padding:7px 14px;border-radius:8px;font-size:.82rem}
.review-panel button.rv-comment-btn{background:transparent;color:var(--txt);border-color:rgba(255,255,255,.14)}
.review-panel button.sm{padding:4px 10px;font-size:.72rem;border-radius:6px}
.review-panel button.rv-c-del{background:transparent;color:#fda4af;border-color:#f4547066}
.review-panel button.rv-c-edit{background:transparent;color:var(--txt);border-color:rgba(255,255,255,.14)}
.review-panel .rv-ai-icon{width:14px;height:14px;vertical-align:-2px;margin-right:6px}
.review-panel .rv-comments{margin-top:10px}
.review-panel .cm.ai{border-color:#a78bfa66}
.rv-verdict{display:inline-block;font-size:.7rem;font-weight:700;letter-spacing:.02em;padding:2px 9px;border-radius:20px;border:1px solid;margin-left:6px;vertical-align:1px}
.review-panel .cm .rv-c-text{margin-top:6px;white-space:pre-wrap}
.review-panel .rv-edit-ta{width:100%;min-height:60px;background:rgba(0,0,0,.28);border:1px solid rgba(255,255,255,.14);border-radius:8px;color:var(--txt);padding:8px;font:inherit;resize:vertical}

/* Custom dropdown: native <select> option lists ignore page CSS on most platforms,
   so the status picker is a fully app-styled div-based listbox instead. */
.csel{position:relative;display:inline-block}
.review-panel .csel-btn{display:flex;align-items:center;gap:10px;cursor:pointer;border:1px solid rgba(255,255,255,.14);background:rgba(0,0,0,.35);color:var(--txt);font:inherit;font-weight:400;padding:8px 12px;border-radius:8px}
.csel-caret{width:10px;height:6px;color:var(--muted)}
.csel-list{position:absolute;top:calc(100% + 6px);left:0;min-width:180px;background:#11162a;border:1px solid rgba(255,255,255,.14);border-radius:10px;box-shadow:0 12px 30px rgba(0,0,0,.55);z-index:20;overflow:hidden;padding:4px}
.csel-list.hide{display:none}
.csel-opt{padding:8px 10px;font-size:.86rem;border-radius:6px;cursor:pointer;white-space:nowrap}
.csel-opt:hover{background:rgba(255,255,255,.08)}
.csel-opt.active{color:#8ab4ff;font-weight:700}

.rv-ai-status{display:none;font-size:.72rem;font-weight:700;padding:4px 11px;border-radius:20px;letter-spacing:.02em}
.rv-ai-status.show{display:inline-block}
.rv-ai-status.testing{background:#332a14;color:#e7b549}
.rv-ai-status.success{background:#0f2e1e;color:#4ade80}
.rv-ai-status.failed{background:#331519;color:#f2555f}

.modal-overlay{position:fixed;inset:0;background:rgba(5,7,14,.65);backdrop-filter:blur(6px);-webkit-backdrop-filter:blur(6px);display:flex;align-items:center;justify-content:center;z-index:100;padding:20px}
.modal-overlay.hide{display:none}
.modal-box{max-width:380px;width:100%}
.modal-msg{font-size:.92rem;color:var(--txt);margin-bottom:20px;line-height:1.5}
.modal-actions{display:flex;justify-content:flex-end;gap:10px}
.modal-actions button{cursor:pointer;border:1px solid #8ab4ff;background:#8ab4ff;color:#08101f;font-weight:700;padding:8px 16px;border-radius:8px;font-size:.86rem}
"""

_JS = """
document.addEventListener('DOMContentLoaded',()=>{
  document.querySelectorAll('.bar-fill,.scorebar i').forEach(el=>{const w=el.getAttribute('data-w');if(w!==null){requestAnimationFrame(()=>{setTimeout(()=>{el.style.width=w+'%';},80);});}});
  document.querySelectorAll('.finding .fhead').forEach(h=>{h.addEventListener('click',()=>{const c=h.closest('.finding');c.classList.toggle('open');c.querySelector('.fbody').classList.toggle('collapsed');});});
  const btns=document.querySelectorAll('.fbtn');btns.forEach(b=>b.addEventListener('click',()=>{btns.forEach(x=>x.classList.remove('active'));b.classList.add('active');const s=b.dataset.sev;document.querySelectorAll('.finding').forEach(f=>{f.style.display=(s==='all'||f.dataset.severity===s)?'':'none';});}));
});

// Finding review: status / comments / AI review — only active when this report is served live
// by the Kodand web UI (pid+sid come from this frame's own URL), not when opened as a static file.
(function(){
  const m = location.pathname.match(/\\/api\\/projects\\/([^/]+)\\/scans\\/([^/]+)\\/report/);
  if(!m) return;
  const pid = m[1];
  const base = `/api/projects/${m[1]}/scans/${m[2]}`;

  // Open File: only meaningful when this report is served live (pid/sid known) —
  // opens files.html on the exact file, highlighted/scrolled to the clicked line.
  document.querySelectorAll('.open-file-btn').forEach(btn=>{
    btn.addEventListener('click', (e)=>{
      e.stopPropagation();
      const file=btn.getAttribute('data-file'), line=btn.getAttribute('data-line')||'1';
      const url=`/static/files.html?project=${encodeURIComponent(pid)}&scan=${encodeURIComponent(m[2])}&path=${encodeURIComponent(file)}&line=${encodeURIComponent(line)}`;
      window.open(url, '_blank');
    });
  });
  const STATUS_LABEL={candidate:'Pending Review','needs-review':'Pending Review',confirmed:'True Positive','false-positive':'False Positive','accepted-risk':'Accepted Risk',remediated:'Remediated',suppressed:'Suppressed'};
  const STATUS_COLOR={candidate:'#94a3b8','needs-review':'#94a3b8',confirmed:'#f43f5e','false-positive':'#22c55e','accepted-risk':'#38bdf8',remediated:'#38bdf8',suppressed:'#94a3b8'};
  const SEV_COLOR={critical:'#f43f5e',high:'#fb923c',medium:'#facc15',low:'#38bdf8',informational:'#94a3b8'};
  const RISK_COLOR={Critical:'#f43f5e',High:'#fb923c',Medium:'#facc15',Low:'#38bdf8'};
  // statuses that mean "no longer an open issue" — excluded from the headline counts
  const RESOLVED=new Set(['false-positive','remediated','suppressed']);
  const esc=s=>(s||'').replace(/[&<>]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;'}[c]));
  function api(url,opts){return fetch(url,opts).then(async r=>{const j=await r.json().catch(()=>({}));if(!r.ok)throw new Error(j.detail||r.statusText);return j;});}

  function appModal(msg,buttons){
    return new Promise(resolve=>{
      const ov=document.getElementById('app-modal');
      document.getElementById('app-modal-msg').textContent=msg;
      const actions=document.getElementById('app-modal-actions');
      actions.innerHTML='';
      buttons.forEach(b=>{
        const btn=document.createElement('button'); btn.type='button'; btn.textContent=b.label;
        if(b.ghost){btn.style.background='transparent';btn.style.color='var(--txt)';btn.style.borderColor='rgba(255,255,255,.14)';}
        if(b.danger){btn.style.background='#f4547022';btn.style.color='#fda4af';btn.style.borderColor='#f4547066';}
        btn.onclick=()=>{ov.classList.add('hide');resolve(b.value);};
        actions.appendChild(btn);
      });
      ov.classList.remove('hide');
    });
  }
  const appConfirm=msg=>appModal(msg,[{label:'Cancel',value:false,ghost:true},{label:'Delete',value:true,danger:true}]);
  const appAlert=msg=>appModal(msg,[{label:'OK',value:true}]);

  function setPill(fid,status){const pill=document.querySelector(`#${CSS.escape(fid)} .rv-status-pill`);if(!pill)return;const c=STATUS_COLOR[status]||'#94a3b8';pill.textContent=STATUS_LABEL[status]||status;pill.style.color=c;pill.style.borderColor=c;pill.style.background=c+'22';}

  // same colours as the review statuses: True positive = confirmed (red), False positive = green
  const VERDICT={true_positive:{label:'True positive',color:'#f43f5e'},false_positive:{label:'False positive',color:'#22c55e'},
                 needs_review:{label:'Needs review',color:'#facc15'}};
  const VERDICT_HEAD=/^[\\s>#*_`\\-]*(?:verdict\\s*[:\\-–—]?\\s*[*_`\\s]*)?(true[\\s-]+positive|false[\\s-]+positive|needs[\\s-]+(?:more[\\s-]+context|review))/i;
  const VERDICT_ANY=/\\bverdict\\s*[:\\-–—]\\s*[*_`\\s]*(true[\\s-]+positive|false[\\s-]+positive|needs[\\s-]+(?:more[\\s-]+context|review))/i;
  function commentVerdict(c){
    if(c.verdict)return c.verdict;
    const head=(c.text||'').slice(0,400), m=head.match(VERDICT_HEAD)||head.match(VERDICT_ANY);
    if(!m)return null;
    const w=m[1].toLowerCase();
    return w.startsWith('true')?'true_positive':w.startsWith('false')?'false_positive':'needs_review';
  }
  function verdictBadge(c){
    const v=VERDICT[commentVerdict(c)]; if(!v)return '';
    return `<span class="rv-verdict" style="color:${v.color};border-color:${v.color}88;background:${v.color}22">${v.label}</span>`;
  }

  function renderComments(panel,comments){
    const el=panel.querySelector('.rv-comments');
    el.innerHTML=(comments&&comments.length)?comments.map(c=>`
      <div class="ev cm ${c.kind==='ai'?'ai':''}" data-cid="${c.id}">
        <div class="row" style="justify-content:space-between;flex-wrap:nowrap">
          <div><b>${esc(c.author)}</b>${verdictBadge(c)} <span class="hash">${esc(c.created)}${c.edited?' (edited)':''}</span></div>
          <div class="row" style="gap:6px">
            <button type="button" class="sm rv-c-edit">Edit</button>
            <button type="button" class="sm rv-c-del">Delete</button>
          </div>
        </div>
        <div class="rv-c-text">${esc(c.text)}</div>
      </div>`).join(''):'<div class="hash">No comments yet.</div>';
    el.querySelectorAll('.rv-c-edit').forEach(btn=>btn.addEventListener('click', ()=>startEditComment(panel, btn.closest('.cm'))));
    el.querySelectorAll('.rv-c-del').forEach(btn=>btn.addEventListener('click', ()=>deleteComment(panel, btn.closest('.cm').dataset.cid)));
  }
  async function refreshComments(panel){const fid=panel.dataset.fid;const rv=await api(base+'/review');renderComments(panel,rv[fid].comments);}
  async function deleteComment(panel,cid){
    if(!await appConfirm('Delete this comment?'))return;
    const fid=panel.dataset.fid;
    try{await api(`${base}/findings/${fid}/comments/${cid}`,{method:'DELETE'});await refreshComments(panel);}
    catch(err){appAlert(err.message);}
  }
  function startEditComment(panel,cmEl){
    const cid=cmEl.dataset.cid, textEl=cmEl.querySelector('.rv-c-text'), current=textEl.textContent;
    textEl.innerHTML=`<textarea class="rv-edit-ta">${esc(current)}</textarea><div class="row" style="margin-top:6px;gap:6px"><button type="button" class="sm rv-edit-save">Save</button><button type="button" class="sm rv-edit-cancel">Cancel</button></div>`;
    cmEl.querySelector('.rv-edit-save').addEventListener('click', async ()=>{
      const val=cmEl.querySelector('.rv-edit-ta').value.trim(); if(!val)return;
      try{await api(`${base}/findings/${panel.dataset.fid}/comments/${cid}`,{method:'PUT',headers:{'Content-Type':'application/json'},body:JSON.stringify({text:val})});await refreshComments(panel);}
      catch(err){appAlert(err.message);}
    });
    cmEl.querySelector('.rv-edit-cancel').addEventListener('click', ()=>refreshComments(panel));
  }

  // Open lists are "portaled" to <body> with position:fixed while open, computed from the
  // button's own bounding box. Each .finding card has backdrop-filter (via .glass), which
  // creates its own stacking context — an absolutely-positioned list nested inside one card
  // still gets painted over by a later sibling card regardless of z-index, so keeping the
  // list inside its card can never fix the overlap; it has to leave that stacking context.
  let openCsel=null;
  function closeOpenDropdown(){
    if(!openCsel)return;
    const {csel,list}=openCsel;
    list.classList.add('hide');
    list.style.position=''; list.style.top=''; list.style.left=''; list.style.minWidth='';
    csel.appendChild(list);
    openCsel=null;
  }
  function initDropdown(panel,initialStatus){
    const csel=panel.querySelector('.rv-status');
    const btn=csel.querySelector('.csel-btn'), list=csel.querySelector('.csel-list');
    const opts=Array.from(list.querySelectorAll('.csel-opt'));
    const apply=v=>{
      csel.dataset.value=v;
      const opt=opts.find(o=>o.dataset.value===v);
      btn.firstChild.textContent=(opt?opt.textContent:STATUS_LABEL[v]||v);
      opts.forEach(o=>o.classList.toggle('active',o.dataset.value===v));
    };
    apply(opts.some(o=>o.dataset.value===initialStatus)?initialStatus:'needs-review');
    btn.addEventListener('click',e=>{
      e.stopPropagation();
      if(openCsel && openCsel.list===list){closeOpenDropdown();return;}
      closeOpenDropdown();
      const r=btn.getBoundingClientRect();
      list.style.position='fixed'; list.style.top=(r.bottom+6)+'px'; list.style.left=r.left+'px'; list.style.minWidth=r.width+'px';
      document.body.appendChild(list);
      list.classList.remove('hide');
      openCsel={csel,list};
    });
    opts.forEach(o=>o.addEventListener('click',async ()=>{
      closeOpenDropdown();
      const v=o.dataset.value; const prev=csel.dataset.value; if(v===prev)return;
      apply(v);
      try{const res=await api(`${base}/findings/${panel.dataset.fid}/status`,{method:'PUT',headers:{'Content-Type':'application/json'},body:JSON.stringify({status:v})});setPill(panel.dataset.fid,res.status);await recomputeSummary();}
      catch(err){apply(prev);appAlert(err.message);}
    }));
  }
  document.addEventListener('click',closeOpenDropdown);
  document.addEventListener('scroll',closeOpenDropdown,true);

  // Claude Agent: the work happens in the user's own Claude Code session (kodand agent), so the
  // button only queues a job; one poller then follows every open job on this page until its
  // verdict comment lands.
  const AGENT_LABEL={queued:'Queued for Claude Agent',running:'Claude Agent verifying…',done:'Complete',failed:'Claude Agent failed',cancelled:'Cancelled'};
  const AGENT_OPEN=['queued','running'];
  const agentWatch=new Map(); let agentTimer=null;
  function showAgentJob(panel,job){
    const badge=panel.querySelector('.rv-ai-status');
    const cls=job.status==='done'?'success':(job.status==='failed'||job.status==='cancelled')?'failed':'testing';
    badge.className='rv-ai-status show '+cls; badge.textContent=AGENT_LABEL[job.status]||job.status; badge.title=job.error||'';
    if(job.status==='done')setTimeout(()=>{badge.className='rv-ai-status';badge.textContent='';badge.title='';},4000);
  }
  async function agentJobsByFinding(){
    const jobs=await api(`/api/agent/jobs?project=${encodeURIComponent(pid)}&scan=${encodeURIComponent(m[2])}`);
    const latest={}; jobs.forEach(j=>{if(!latest[j.finding_id])latest[j.finding_id]=j;});  // newest first
    return latest;
  }
  async function pollAgentJobs(){
    let latest; try{latest=await agentJobsByFinding();}catch(e){return;}
    for(const [fid,panel] of agentWatch){
      const job=latest[fid]; if(!job)continue;
      showAgentJob(panel,job);
      if(!AGENT_OPEN.includes(job.status)){agentWatch.delete(fid); if(job.status==='done')await refreshComments(panel);}
    }
    if(!agentWatch.size){clearInterval(agentTimer);agentTimer=null;}
  }
  function watchAgentJob(fid,panel){agentWatch.set(fid,panel); if(!agentTimer)agentTimer=setInterval(pollAgentJobs,3000);}

  let LAST_REVIEW={};
  async function recomputeSummary(review){
    if(!review){try{review=await api(base+'/review');}catch(e){return;}}
    LAST_REVIEW=review;
    const sevCounts={critical:0,high:0,medium:0,low:0,informational:0};
    const catCounts={};
    let total=0;
    document.querySelectorAll('.finding').forEach(fEl=>{
      const st=(review[fEl.id]||{}).status||'candidate';
      const resolved=RESOLVED.has(st);
      fEl.classList.toggle('resolved',resolved);
      if(resolved)return;
      total++;
      const sev=fEl.dataset.severity; if(sevCounts[sev]!==undefined)sevCounts[sev]++;
      const cat=fEl.dataset.category; catCounts[cat]=(catCounts[cat]||0)+1;
    });
    const fc=document.getElementById('findings-count'); if(fc)fc.textContent=total;
    Object.entries(sevCounts).forEach(([k,v])=>{const el=document.querySelector(`.stat[data-sev="${k}"] .n`);if(el)el.textContent=v;});
    let risk='Informational';
    for(const k of ['critical','high','medium','low']){if(sevCounts[k]>0){risk=k.charAt(0).toUpperCase()+k.slice(1);break;}}
    const rv=document.getElementById('risk-value'); if(rv){rv.textContent=risk;rv.style.color=RISK_COLOR[risk]||'#94a3b8';}
    const rd=document.getElementById('risk-dot'); if(rd)rd.style.background=RISK_COLOR[risk]||'#94a3b8';
    const totalSev=Object.values(sevCounts).reduce((a,b)=>a+b,0)||1;
    let acc=0; const stops=[];
    ['critical','high','medium','low','informational'].forEach(k=>{
      const c=sevCounts[k], frac=c/totalSev, start=acc*360; acc+=frac; const end=acc*360;
      if(c>0)stops.push(`${SEV_COLOR[k]} ${start.toFixed(2)}deg ${end.toFixed(2)}deg`);
      const val=document.querySelector(`.legend .row[data-sev="${k}"] .val`); if(val)val.textContent=c;
    });
    const donut=document.getElementById('donut'); if(donut)donut.style.setProperty('--gc',stops.length?stops.join(', '):'#334155 0deg 360deg');
    const dt=document.getElementById('donut-total'); if(dt)dt.textContent=total;
    const maxCnt=Math.max(1,...Object.values(catCounts));
    document.querySelectorAll('.bar-row').forEach(row=>{
      const cnt=catCounts[row.dataset.cat]||0, w=(cnt/maxCnt*100);
      const fill=row.querySelector('.bar-fill'); if(fill){fill.style.width=w+'%';fill.setAttribute('data-w',w.toFixed(1));}
      const cntEl=row.querySelector('.cnt'); if(cntEl)cntEl.textContent=cnt;
    });
  }

  document.addEventListener('DOMContentLoaded', async ()=>{
    const wrap=document.querySelector('.wrap');
    if(wrap){
      const nav=document.createElement('div');
      nav.className='report-nav';
      nav.innerHTML=`<a href="/">&larr; Home</a><a href="/#p=${encodeURIComponent(pid)}">&larr; Back to Scans</a>`;
      wrap.insertBefore(nav, wrap.firstChild);
    }
    let review={};
    try{review=await api(base+'/review');}catch(e){return;}
    LAST_REVIEW=review;
    document.querySelectorAll('.review-panel').forEach(panel=>{
      const fid=panel.dataset.fid, r=review[fid]||{status:'needs-review',comments:[]};
      initDropdown(panel,r.status);
      renderComments(panel,r.comments); setPill(fid,r.status);
      panel.querySelector('.rv-comment-btn').addEventListener('click', async ()=>{
        const input=panel.querySelector('.rv-comment-input'), text=input.value.trim(); if(!text)return;
        try{await api(`${base}/findings/${fid}/comments`,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({text})});input.value='';await refreshComments(panel);}
        catch(err){appAlert(err.message);}
      });
      panel.querySelector('.rv-ai-btn').addEventListener('click', async ()=>{
        const badge=panel.querySelector('.rv-ai-status'); badge.className='rv-ai-status show testing'; badge.textContent='Reviewing…';
        try{
          await api(`${base}/findings/${fid}/ai-review`,{method:'POST'});
          await refreshComments(panel);
          badge.className='rv-ai-status show success'; badge.textContent='Complete';
          setTimeout(()=>{badge.className='rv-ai-status';badge.textContent='';},4000);
        }catch(err){
          badge.className='rv-ai-status show failed'; badge.textContent='Failed';
          appAlert(err.message);
        }
      });
      panel.querySelector('.rv-claude-agent-btn').addEventListener('click', async ()=>{
        try{
          const job=await api(`${base}/findings/${fid}/agent-verify`,{method:'POST'});
          showAgentJob(panel,job); watchAgentJob(fid,panel);
          const s=await api('/api/agent/status');
          if(s.state==='offline')appAlert('Queued. Claude Agent isn\\'t running yet — start it with `kodand agent` in a terminal and this finding will be picked up as soon as it starts.');
        }catch(err){appAlert(err.message);}
      });
      panel.querySelector('.rv-chat-btn').addEventListener('click', async ()=>{
        const input=panel.querySelector('.rv-chat-input'), text=input.value.trim(); if(!text)return;
        const badge=panel.querySelector('.rv-ai-status'); badge.className='rv-ai-status show testing'; badge.textContent='Thinking…';
        try{
          await api(`${base}/findings/${fid}/ai-chat`,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({message:text})});
          input.value=''; await refreshComments(panel);
          badge.className='rv-ai-status show success'; badge.textContent='Complete';
          setTimeout(()=>{badge.className='rv-ai-status';badge.textContent='';},4000);
        }catch(err){
          await refreshComments(panel); // the question is saved even if the AI call failed
          badge.className='rv-ai-status show failed'; badge.textContent='Failed';
          appAlert(err.message);
        }
      });
    });
    recomputeSummary(review);
    // pick up jobs queued earlier (e.g. by auto-verify on this scan) so their progress shows too
    try{
      const latest=await agentJobsByFinding();
      document.querySelectorAll('.review-panel').forEach(panel=>{
        const job=latest[panel.dataset.fid];
        if(job&&AGENT_OPEN.includes(job.status)){showAgentJob(panel,job);watchAgentJob(panel.dataset.fid,panel);}
      });
    }catch(e){}
  });
})();
"""


def _pill(text, color):
    return f'<span class="pill" style="color:{color};border-color:{color};background:{color}22">{html.escape(text)}</span>'


def _score_color(s):
    return "#22c55e" if s >= 70 else "#facc15" if s >= 45 else "#94a3b8"


def write_html(a, p):
    p = Path(p); p.parent.mkdir(parents=True, exist_ok=True)
    s = a.summary
    sev_order = ["critical", "high", "medium", "low", "informational"]
    sev_counts = {k: getattr(s, k) for k in sev_order}
    total_sev = sum(sev_counts.values()) or 1
    stops, acc, legend_rows = [], 0.0, ""
    for k in sev_order:
        c = sev_counts[k]; frac = c / total_sev; start = acc * 360; acc += frac; end = acc * 360; col = SEV_COLORS[k]
        if c > 0:
            stops.append(f"{col} {start:.2f}deg {end:.2f}deg")
        legend_rows += f'<div class="row" data-sev="{k}"><span class="swatch" style="background:{col}"></span>{k.title()}<span class="val">{c}</span></div>'
    gc = ", ".join(stops) if stops else "#334155 0deg 360deg"
    stat_cards = ""
    for k in sev_order:
        stat_cards += f'<div class="stat" data-sev="{k}" style="--accent:{SEV_COLORS[k]}"><div class="n" style="color:{SEV_COLORS[k]}">{sev_counts[k]}</div><div class="l">{k}</div></div>'
    cov = {x.category: x for x in a.coverage}
    max_cnt = max([x.finding_count for x in a.coverage] + [1])
    status_color = {"reviewed-findings": "#fb923c", "reviewed-no-findings": "#22c55e", "partially-reviewed": "#facc15", "not-supported": "#94a3b8", "not-applicable": "#94a3b8"}

    def _taxonomy_tables(prefix):
        """AST/MCP/HOOK are three independent, non-overlapping taxonomies (see
        kodand.categories) — each gets its own bar chart + coverage matrix so a
        reader never has to mentally separate an AST code from an MCP/HOOK one
        in the same list."""
        order = TAXONOMIES[prefix]["order"]
        extra = [c for c in cov if _taxonomy_of(c) == prefix and c not in order]
        cats = order + extra
        bars = ""
        for cat in cats:
            item = cov.get(cat); name = item.name if item else ""; cnt = item.finding_count if item else 0
            col = CATEGORY_COLORS.get(cat, "#888"); w = (cnt / max_cnt * 100) if max_cnt else 0
            bars += f'<div class="bar-row" data-cat="{cat}" title="{html.escape(name)}"><span class="cat" style="color:{col}">{cat}</span><span class="bar-track"><span class="bar-fill" data-w="{w:.1f}" style="background:linear-gradient(90deg,{col}88,{col})"></span></span><span class="cnt">{cnt}</span></div>'
        rows = ""
        for cat in cats:
            x = cov.get(cat)
            if x is None:
                continue
            col = CATEGORY_COLORS.get(x.category, "#888")
            rows += f"<tr><td><span class='dot' style='display:inline-block;background:{col}'></span> <b style='color:{col}'>{x.category}</b></td><td>{html.escape(x.name)}</td><td>{_pill(x.status.replace('-', ' ').title(), status_color.get(x.status, '#94a3b8'))}</td><td>{x.finding_count}</td></tr>"
        if not rows:
            rows = "<tr><td colspan='4'>No categories in this taxonomy were evaluated</td></tr>"
        return bars, rows

    ast_bar_rows, ast_coverage_rows = _taxonomy_tables("AST")
    mcp_bar_rows, mcp_coverage_rows = _taxonomy_tables("MCP")
    hook_bar_rows, hook_coverage_rows = _taxonomy_tables("HOOK")
    plats = {x.platform: x for x in a.inventory.platforms}
    platform_rows = ""
    for k, v in plats.items():
        ev = "".join(f"<li><code>{html.escape(x)}</code></li>" for x in v.evidence)
        platform_rows += f"<tr><td><b>{html.escape(k)}</b></td><td>{_pill('Detected', '#22c55e')}</td><td>{_pill(v.confidence.title(), '#38bdf8')}</td><td><ul>{ev}</ul></td></tr>"
    if not platform_rows:
        platform_rows = "<tr><td colspan='4'>No platform indicators detected</td></tr>"
    rel_rows = ""
    for ext in a.inventory.extensions:
        und = "".join(f'<span class="tag" style="background:#facc1522;color:#facc15">{html.escape(x)}</span>' for x in ext.undeclared_supporting_files) or _pill("None", "#22c55e")
        brk = "".join(f'<span class="tag" style="background:#f43f5e22;color:#fda4af">{html.escape(x)}</span>' for x in ext.broken_references) or _pill("None", "#22c55e")
        rel_rows += f"<tr><td><code>{html.escape(ext.entrypoint)}</code></td><td>{und}</td><td>{brk}</td></tr>"
    if not rel_rows:
        rel_rows = "<tr><td colspan='3'>No skills discovered</td></tr>"
    def _open_file_btn(file_path, line, label="Open File ↗"):
        # wired up client-side only when this report is served live (see _JS) —
        # opens files.html on the exact file, scrolled/highlighted to this line
        return (f'<button type="button" class="open-file-btn" data-file="{html.escape(file_path, quote=True)}" '
                f'data-line="{line or 1}" title="Open this file, highlighted at line {line or 1}">{label}</button>')

    def _finding_html(f):
        sev_col = SEV_COLORS.get(f.severity, "#888"); cat_col = CATEGORY_COLORS.get(f.category, "#888")
        occ = ""
        for o in f.occurrences:
            kindtag = f'<span class="tag">{o.match_kind}</span>' + (f'<span class="tag">decoded:{html.escape(o.decoded_from)}</span>' if getattr(o, "decoded_from", None) else "") + (f'<span class="tag">{html.escape(o.context)}</span>' if getattr(o, "context", None) else "")
            occ += (f'<div class="ev"><code>{html.escape(o.file)}:{o.line}:{o.column_start or 1}</code> {kindtag} '
                    f'{_open_file_btn(o.file, o.line, "Open ↗")}<br><b>Matched:</b> <code>{html.escape(o.matched_value)}</code>'
                    f'<br>{html.escape(o.excerpt)}<br><span class="hash">Evidence hash: {o.evidence_hash}</span></div>')
        primary = f.occurrences[0] if f.occurrences else None
        primary_btn = _open_file_btn(primary.file, primary.line) if primary else ""
        files = "".join(f"<li><code>{html.escape(x)}</code></li>" for x in f.affected_files)
        caps = ", ".join(map(html.escape, f.capabilities)) or "None inferred"
        sc = f.confidence_score; sc_col = _score_color(sc)
        merged = f'<span class="tag">merged: {html.escape(", ".join(f.merged_rule_ids))}</span>' if f.merged_rule_ids else ""
        prox = "".join(f'<span class="tag" style="background:#f43f5e22;color:#fda4af">⚡ {html.escape(x)}</span>' for x in f.proximity_signals)
        rel_sig = "".join(f'<span class="tag" style="background:#facc1522;color:#facc15">🔗 {html.escape(x)}</span>' for x in getattr(f, "relationship_signals", []))
        return f"""
<div class="glass finding" id="{html.escape(f.finding_id)}" data-severity="{html.escape(f.severity)}" data-category="{html.escape(f.category)}" style="--sev:{sev_col}">
  <div class="fhead">
    <div>
      <div class="ftitle">{html.escape(f.title)} {prox}{rel_sig}</div>
      <div class="fmeta">{html.escape(f.finding_id)} &middot; {html.escape(f.rule_id)} &middot; <span style="color:{cat_col}">{html.escape(f.category)} {html.escape(f.category_name)}</span> &middot; {f.occurrence_count} occurrence(s) {merged} {primary_btn}</div>
    </div>
    <div class="badges">
      {_pill(f.severity.upper(), sev_col)}
      <span class="pill" style="color:{sc_col};border-color:{sc_col};background:{sc_col}22">SCORE {sc}</span>
      {_pill(f.confidence.upper()+' CONF', '#38bdf8')}
      <span class="pill rv-status-pill" style="color:#94a3b8;border-color:#94a3b8;background:#94a3b822">{html.escape(f.status.replace('-',' ').title())}</span>
    </div>
    <span class="caret">&#9656;</span>
  </div>
  <div class="fbody collapsed">
    <div><h4>Confidence Score</h4><div class="scorebar"><i data-w="{sc}" style="width:0;background:{sc_col}"></i></div><p class="tagline">{sc}/100 composite (context · evidence · corroboration · proximity · relationship)</p></div>
    <div><h4>Capabilities</h4><p>{caps}</p></div>
    <div><h4>Description</h4><p>{html.escape(f.description)}</p></div>
    <div><h4>Technical Root Cause</h4><p>{html.escape(f.technical_root_cause)}</p></div>
    <div><h4>Attack Scenario</h4><p>{html.escape(f.attack_scenario)}</p></div>
    <div><h4>Potential Impact</h4><p>{html.escape(f.potential_impact)}</p></div>
    <div class="full"><h4>Affected Files</h4><ul>{files}</ul></div>
    <div class="full"><h4>Evidence &amp; Locations</h4>{occ}</div>
    <div class="full"><h4>Recommendation</h4><p>{html.escape(f.remediation)}</p></div>
    <div class="full"><h4>Validation Guidance</h4><p>{html.escape(f.validation_guidance)}</p></div>
    <div class="full review-panel" data-fid="{html.escape(f.finding_id)}">
      <h4>Review</h4>
      <div class="row">
        <div class="csel rv-status" tabindex="0" data-value="needs-review">
          <button type="button" class="csel-btn">Pending Review<svg viewBox="0 0 10 6" class="csel-caret"><path d="M0 0l5 6 5-6z" fill="currentColor"/></svg></button>
          <div class="csel-list hide">
            <div class="csel-opt" data-value="needs-review">Pending Review</div>
            <div class="csel-opt" data-value="confirmed">True Positive</div>
            <div class="csel-opt" data-value="false-positive">False Positive</div>
            <div class="csel-opt" data-value="remediated">Remediated</div>
            <div class="csel-opt" data-value="suppressed">Suppress</div>
          </div>
        </div>
        <button type="button" class="rv-ai-btn" title="Your AI provider (Settings → AI assistant) reviews the finding with the full affected files, and opens other scanned files if it needs to"><svg class="rv-ai-icon" viewBox="0 0 24 24" fill="currentColor"><path d="M12 2.5l1.7 5 5 1.7-5 1.7-1.7 5-1.7-5-5-1.7 5-1.7 1.7-5z"/><path d="M19.3 13l.9 2.6 2.6.9-2.6.9-.9 2.6-.9-2.6-2.6-.9 2.6-.9.9-2.6z"/></svg>AI Review</button>
        <button type="button" class="rv-ai-btn rv-claude-agent-btn" title="Send this finding to your Claude Agent session (kodand agent) — it verifies with your own Claude login and adds the verdict as a comment"><svg class="rv-ai-icon" viewBox="0 0 24 24" fill="currentColor"><path d="M12 2.5l1.7 5 5 1.7-5 1.7-1.7 5-1.7-5-5-1.7 5-1.7 1.7-5z"/><path d="M19.3 13l.9 2.6 2.6.9-2.6.9-.9 2.6-.9-2.6-2.6-.9 2.6-.9.9-2.6z"/></svg>Verify with Claude Agent</button>
        <span class="rv-ai-status"></span>
      </div>
      <div class="rv-comments"></div>
      <div class="row" style="margin-top:8px">
        <input type="text" class="rv-comment-input" placeholder="Add a comment…">
        <button type="button" class="rv-comment-btn">Comment</button>
      </div>
      <div class="row" style="margin-top:8px">
        <input type="text" class="rv-chat-input" placeholder="Ask AI a follow-up about this finding…">
        <button type="button" class="rv-chat-btn"><svg class="rv-ai-icon" viewBox="0 0 24 24" fill="currentColor"><path d="M12 2.5l1.7 5 5 1.7-5 1.7-1.7 5-1.7-5-5-1.7 5-1.7 1.7-5z"/><path d="M19.3 13l.9 2.6 2.6.9-2.6.9-.9 2.6-.9-2.6-2.6-.9 2.6-.9.9-2.6z"/></svg>Ask AI</button>
      </div>
    </div>
  </div>
</div>"""
    # Group findings into their taxonomy (AST/MCP/HOOK) so each gets its own
    # section — a reader looking for "what's wrong with my MCP servers" should
    # never have to pick AST findings out of the same list.
    findings_by_tax = {"AST": [], "MCP": [], "HOOK": [], "OTHER": []}
    for f in a.findings:
        findings_by_tax[_taxonomy_of(f.category) or "OTHER"].append(f)

    def _toc_items(items):
        return "".join(f'<li><a href="#{html.escape(f.finding_id)}">{_pill(f.severity.upper(), SEV_COLORS.get(f.severity,"#888"))} {html.escape(f.title)} <span class="tagline">({f.confidence_score})</span></a></li>' for f in items)

    def _findings_section(title, items, empty_note):
        if not items:
            return f'<div class="section-title">{html.escape(title)}</div><div class="glass"><p class="tagline">{html.escape(empty_note)}</p></div>'
        body = "".join(_finding_html(f) for f in items)
        return f'<div class="section-title">{html.escape(title)}</div>{body}'

    _NO_FINDINGS_LI = '<li class="tagline">None</li>'
    toc_sections = ""
    for label, key in (("Agentic Skills Findings (AST Top 10)", "AST"),
                       ("MCP Server Findings (OWASP MCP Top 10)", "MCP"),
                       ("Hook Findings (Hooks Top 10)", "HOOK")):
        items = findings_by_tax[key]
        toc_sections += f'<div class="toc-group"><h4>{html.escape(label)} ({len(items)})</h4><ul class="toc">{_toc_items(items) or _NO_FINDINGS_LI}</ul></div>'
    if findings_by_tax["OTHER"]:
        toc_sections += f'<div class="toc-group"><h4>Other Findings ({len(findings_by_tax["OTHER"])})</h4><ul class="toc">{_toc_items(findings_by_tax["OTHER"])}</ul></div>'

    ast_findings_html = _findings_section("Agentic Skills Findings (AST Top 10)", findings_by_tax["AST"],
                                          "No Agentic Skills Top 10 findings.")
    mcp_findings_html = _findings_section("MCP Server Findings (OWASP MCP Top 10)", findings_by_tax["MCP"],
                                          "No MCP server config findings — scan a repository with an .mcp.json/claude_desktop_config.json to see this section populated.")
    hook_findings_html = _findings_section("Hook Findings (Hooks Top 10)", findings_by_tax["HOOK"],
                                           "No hook findings — scan a repository with a .claude/settings.json to see this section populated.")
    other_findings_html = _findings_section("Other Findings", findings_by_tax["OTHER"], "") if findings_by_tax["OTHER"] else ""
    limitations = "".join(f"<li>{html.escape(x)}</li>" for x in a.limitations)
    risk_col = {"Critical": "#f43f5e", "High": "#fb923c", "Medium": "#facc15", "Low": "#38bdf8"}.get(a.summary.overall_risk, "#94a3b8")
    supp = f'<span class="chip">Suppressed <b>{a.summary.suppressed}</b></span>' if getattr(a.summary, "suppressed", 0) else ""
    corr_risks = getattr(a, "correlated_risks", None)
    corr_chip = f'<span class="chip">Correlated Risks <b>{html.escape(", ".join(corr_risks))}</b></span>' if corr_risks else ""
    doc = f"""<!doctype html><html lang="en"><head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Kodand Security Assessment — {html.escape(a.inventory.repository_name)}</title>
<link rel="icon" href="{_FAVICON_DATA_URI}">
<style>{_CSS}</style></head>
<body id="top"><div class="wrap">
  <div class="glass hero"><div class="brand"><img src="{_LOGO_DATA_URI}" alt="Kodand — Agentic AI Skills Threat and Risk Assessment" class="logo-lockup"><div class="brand-label">Security Assessment</div></div>
    <div class="chips">
      <span class="chip">Repository <b>{html.escape(a.inventory.repository_name)}</b></span>
      <span class="chip"><span class="dot" id="risk-dot" style="background:{risk_col}"></span>Overall Risk <b id="risk-value" style="color:{risk_col}">{a.summary.overall_risk}</b></span>
      <span class="chip">Findings <b id="findings-count">{a.summary.total}</b></span>{supp}{corr_chip}
      <span class="chip">Skills <b>{len(a.inventory.extensions)}</b></span>
      <span class="chip">Files <b>{a.inventory.files_examined}</b></span>
      <span class="chip">Kodand <b>{a.scanner_version}</b></span>
      <span class="chip">Rule Pack <b>{a.rule_pack_version}</b></span>
      <span class="chip">{a.analyzed_at.strftime('%Y-%m-%d %H:%M UTC')}</span>
    </div></div>
  <div class="section-title">Severity Overview</div><div class="glass"><div class="stats">{stat_cards}</div></div>
  <div class="section-title">Risk Distribution</div>
  <div class="glass"><div class="charts">
    <div class="donut-wrap"><div class="donut" id="donut" style="--gc:{gc}"><div class="donut-center"><span class="big" id="donut-total">{a.summary.total}</span><span class="sm">Findings</span></div></div><div class="legend">{legend_rows}</div></div>
    <div class="bars-cols">
      <div class="bars"><div class="section-title" style="margin:0 0 4px">AST Top 10</div>{ast_bar_rows}</div>
      <div class="bars"><div class="section-title" style="margin:0 0 4px">MCP Top 10</div>{mcp_bar_rows}</div>
      <div class="bars"><div class="section-title" style="margin:0 0 4px">Hooks Top 10</div>{hook_bar_rows}</div>
    </div>
  </div></div>
  <div class="section-title">Detected AI Platforms</div>
  <div class="glass"><table><thead><tr><th>Platform</th><th>Status</th><th>Confidence</th><th>Detection Evidence</th></tr></thead><tbody>{platform_rows}</tbody></table>
    <p class="tagline" style="margin-top:12px">Fingerprinting indicates matching repository artifacts and does not independently confirm active deployment.</p></div>
  <div class="section-title">Skill File Relationships</div>
  <div class="glass"><table><thead><tr><th>Entrypoint</th><th>Undeclared Supporting Files</th><th>Broken References</th></tr></thead><tbody>{rel_rows}</tbody></table>
    <p class="tagline" style="margin-top:12px">Undeclared: exists on disk but never named in the skill's own instructions. Broken: named in the instructions but not found on disk. Candidate signals for human review.</p></div>
  <div class="section-title">Assessment Scope &amp; Methodology</div>
  <div class="glass"><p>Repository inventory &rarr; platform fingerprinting &rarr; skill discovery &rarr; literal+regex matching &rarr; encoding normalization &rarr; context v2 &rarr; capability inference &rarr; dedup &amp; proximity &rarr; composite scoring &rarr; suppression &rarr; report.</p>
    <p><b>Repository root:</b> <code>{html.escape(a.inventory.repository_root)}</code></p>
    <p><b>Excluded paths:</b> {', '.join(map(html.escape, a.inventory.excluded_paths)) or 'None'} &nbsp;&middot;&nbsp; <b>Warnings:</b> {len(a.inventory.warnings)}</p></div>
  <div class="section-title">AST Top 10 Coverage Matrix</div>
  <div class="glass"><table><thead><tr><th>Category</th><th>Name</th><th>Status</th><th>Findings</th></tr></thead><tbody>{ast_coverage_rows}</tbody></table></div>
  <div class="section-title">MCP Top 10 Coverage Matrix <span class="tagline">(OWASP MCP Top 10, 2025)</span></div>
  <div class="glass"><table><thead><tr><th>Category</th><th>Name</th><th>Status</th><th>Findings</th></tr></thead><tbody>{mcp_coverage_rows}</tbody></table></div>
  <div class="section-title">Hooks Top 10 Coverage Matrix</div>
  <div class="glass"><table><thead><tr><th>Category</th><th>Name</th><th>Status</th><th>Findings</th></tr></thead><tbody>{hook_coverage_rows}</tbody></table></div>
  <div class="section-title">Table of Contents</div><div class="glass toc-wrap">{toc_sections}</div>
  <div class="filters"><span class="fbtn active" data-sev="all">All ({a.summary.total})</span><span class="fbtn" data-sev="critical">Critical ({s.critical})</span><span class="fbtn" data-sev="high">High ({s.high})</span><span class="fbtn" data-sev="medium">Medium ({s.medium})</span><span class="fbtn" data-sev="low">Low ({s.low})</span><span class="fbtn" data-sev="informational">Info ({s.informational})</span></div>
  {ast_findings_html}
  {mcp_findings_html}
  {hook_findings_html}
  {other_findings_html}
  <div class="section-title">Limitations</div><div class="glass"><ul>{limitations}</ul></div>
  <div class="foot">Generated by Kodand {a.scanner_version} &middot; {a.rule_pack_version} &middot; Candidate findings require human validation.</div>
</div><a href="#top" class="top-link">&#8679; Top</a>
<div class="modal-overlay hide" id="app-modal"><div class="glass modal-box">
  <div class="modal-msg" id="app-modal-msg"></div>
  <div class="modal-actions" id="app-modal-actions"></div>
</div></div>
<script>{_JS}</script></body></html>"""
    p.write_text(doc, encoding='utf-8')


def write_sarif(a, p):
    p = Path(p); p.parent.mkdir(parents=True, exist_ok=True)
    rules, results = {}, []
    for f in a.findings:
        rules[f.rule_id] = {"id": f.rule_id, "name": f.title, "shortDescription": {"text": f.description}, "help": {"text": f.remediation}}
        for o in f.occurrences:
            results.append({"ruleId": f.rule_id, "level": "error" if f.severity in {"critical", "high"} else "warning" if f.severity == "medium" else "note", "message": {"text": f.title}, "locations": [{"physicalLocation": {"artifactLocation": {"uri": o.file}, "region": {"startLine": o.line, "startColumn": o.column_start or 1}}}], "properties": {"findingId": f.finding_id, "confidence": f.confidence, "confidenceScore": f.confidence_score, "status": f.status, "matchKind": o.match_kind}})
    p.write_text(json.dumps({"version": "2.1.0", "$schema": "https://json.schemastore.org/sarif-2.1.0.json", "runs": [{"tool": {"driver": {"name": "Kodand", "version": a.scanner_version, "rules": list(rules.values())}}, "results": results}]}, indent=2), encoding='utf-8')


def write_docx(a, p):
    d = Document(); sec = d.sections[0]; sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7)
    d.add_picture(str(_ICON_PATH), width=Inches(.6))
    d.add_heading('Kodand Security Assessment', 0)
    d.add_paragraph(f"Repository: {a.inventory.repository_name}")
    d.add_paragraph(f"Generated: {a.analyzed_at.isoformat()} | Kodand {a.scanner_version} | {a.rule_pack_version}")
    d.add_paragraph('Candidate findings require human validation.')
    d.add_heading('Executive Summary', 1)
    d.add_paragraph(f"Overall risk: {a.summary.overall_risk}. Correlated findings: {a.summary.total} (suppressed {getattr(a.summary,'suppressed',0)}). Critical {a.summary.critical}, High {a.summary.high}, Medium {a.summary.medium}, Low {a.summary.low}.")
    d.add_heading('Detected AI Platforms', 1)
    for x in a.inventory.platforms:
        d.add_paragraph(f"{x.platform} ({x.confidence} confidence): {', '.join(x.evidence)}", style='List Bullet')
    d.add_heading('AST10 Coverage Matrix', 1)
    for x in a.coverage:
        d.add_paragraph(f"{x.category} {x.name}: {x.status}, {x.finding_count} finding(s)", style='List Bullet')
    for f in a.findings:
        d.add_page_break()
        d.add_heading(f"{f.finding_id}: {f.title}", 1)
        d.add_paragraph(f"{f.category} {f.category_name} | Severity: {f.severity.title()} | Confidence: {f.confidence.title()} (score {f.confidence_score}) | Status: {f.status.title()}")
        if f.proximity_signals:
            d.add_paragraph("Proximity signals: " + ", ".join(f.proximity_signals))
        d.add_heading('Description', 2); d.add_paragraph(f.description)
        d.add_heading('Technical Root Cause', 2); d.add_paragraph(f.technical_root_cause)
        d.add_heading('Attack Scenario and Impact', 2); d.add_paragraph(f.attack_scenario); d.add_paragraph(f.potential_impact)
        d.add_heading(f"Evidence ({f.occurrence_count} occurrences)", 2)
        for o in f.occurrences:
            d.add_paragraph(f"{o.file}:{o.line}:{o.column_start or 1} [{o.match_kind}]\n{o.excerpt}", style='List Bullet')
        d.add_heading('Recommendation', 2); d.add_paragraph(f.remediation)
        d.add_heading('Validation Guidance', 2); d.add_paragraph(f.validation_guidance)
    d.save(p)


def write_pdf(a, p):
    styles = getSampleStyleSheet()
    story = [Image(str(_ICON_PATH), width=40, height=40), Paragraph('Kodand Security Assessment', styles['Title']), Paragraph(f'Repository: {a.inventory.repository_name}', styles['Normal']), Paragraph(f'Overall Risk: {a.summary.overall_risk} | Findings: {a.summary.total}', styles['Heading2']), Spacer(1, 12)]
    for f in a.findings:
        story += [Paragraph(f'{html.escape(f.finding_id)}: {html.escape(f.title)}', styles['Heading1']), Paragraph(f'{html.escape(f.category)} {html.escape(f.category_name)} | {f.severity.title()} | Confidence {f.confidence.title()} (score {f.confidence_score})', styles['Normal']), Paragraph(html.escape(f.description), styles['Normal']), Paragraph('Evidence', styles['Heading2'])]
        for o in f.occurrences:
            story += [Preformatted(f'{o.file}:{o.line}:{o.column_start or 1} [{o.match_kind}]\n{o.excerpt}', styles['Code'])]
        story += [Paragraph('Recommendation', styles['Heading2']), Paragraph(html.escape(f.remediation), styles['Normal']), Spacer(1, 14)]
    SimpleDocTemplate(str(p), pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36).build(story)
